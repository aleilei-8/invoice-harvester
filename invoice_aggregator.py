#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
拼多多开票信息汇总器
====================
把「拼多多开票信息抓取器」插件导出的多个 CSV，或平台自带导出的订单/发票 Excel，
按【订单号】合并去重、补全快递单号、计算金额合计，生成带格式的 .xlsx 报表。

用法：
  python invoice_aggregator.py -i <输入目录或文件> [-o 输出.xlsx]

示例：
  python invoice_aggregator.py -i ./抓取导出 -o 开票汇总.xlsx
  python invoice_aggregator.py -i 发票页.csv 订单页.csv -o 开票汇总.xlsx
"""
import argparse
import csv
import glob
import os
import re
import sys
from collections import OrderedDict
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 标准字段： (key, 中文列名)
FIELDS = [
    ("buyer_title", "发票抬头"),
    ("tax_no", "税号"),
    ("order_id", "订单号"),
    ("tracking_no", "快递单号"),
    ("amount", "金额"),
    ("items", "物品"),
    ("invoice_date", "开票时间"),
    ("apply_date", "申请时间"),
    ("buyer", "买家"),
]

# 各字段可能的表头别名（兼容插件导出 & 平台自带导出）
COL_ALIASES = {
    "buyer_title": ["发票抬头", "抬头", "公司名称", "购买方名称", "购方名称"],
    "tax_no": ["税号", "纳税人识别号", "统一社会信用代码", "识别号", "信用代码"],
    "order_id": ["订单号", "订单编号", "订单", "单号"],
    "tracking_no": ["快递单号", "物流单号", "运单号", "货运单号"],
    "amount": ["金额", "价税合计", "开票金额", "实付金额", "合计", "支付金额", "订单实付金额"],
    "items": ["物品", "商品", "发票内容", "货物", "货品", "商品名称", "品名"],
    "invoice_date": ["开票时间", "开票日期"],
    "apply_date": ["申请时间", "申请日期"],
    "buyer": ["买家", "收货人", "购买方", "收件人", "客户"],
}

AMOUNT_IDX = [i for i, (k, _) in enumerate(FIELDS) if k == "amount"][0]


def norm_header(h):
    return (h or "").strip().lstrip("﻿").strip()


def map_columns(headers):
    """返回 {标准key: 列索引}，每个标准字段只取最匹配的一列。"""
    m = {}
    for idx, h in enumerate(headers):
        hn = norm_header(h)
        for key, aliases in COL_ALIASES.items():
            if key in m:
                continue
            if hn in aliases:
                m[key] = idx
                break
    return m


def read_rows(path):
    """读取一个 csv/xlsx，返回 [(colmap, row), ...]。"""
    ext = os.path.splitext(path)[1].lower()
    rows = []
    if ext == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            data = list(csv.reader(f))
        if not data:
            return rows
        colmap = map_columns(data[0])
        for r in data[1:]:
            rows.append((colmap, r))
    elif ext in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
        wb.close()
        if not data:
            return rows
        headers = ["" if x is None else str(x) for x in data[0]]
        colmap = map_columns(headers)
        for r in data[1:]:
            rows.append((colmap, ["" if x is None else str(x) for x in r]))
    else:
        raise ValueError("不支持的文件类型: " + ext)
    return rows


def clean_amount(s):
    if not s:
        return 0.0
    s = str(s).replace("¥", "").replace("￥", "").replace(",", "").replace("，", "").strip()
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else 0.0


def sort_key(rec):
    for k in ("invoice_date", "apply_date"):
        v = rec.get(k, "").strip()
        if v:
            for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d"):
                try:
                    return (0, datetime.strptime(v, fmt))
                except ValueError:
                    continue
    return (1, datetime.min)


def aggregate(files):
    lib = OrderedDict()
    stats = {"files": 0, "rows": 0, "added": 0, "merged": 0}
    for path in files:
        try:
            rows = read_rows(path)
        except Exception as e:
            print(f"  [跳过] {path}: {e}")
            continue
        if not rows:
            print(f"  [跳过] {path}: 无数据行")
            continue
        stats["files"] += 1
        for cm, row in rows:
            stats["rows"] += 1
            rec = {}
            for key, _ in FIELDS:
                ci = cm.get(key)
                rec[key] = row[ci].strip() if (ci is not None and ci < len(row)) else ""
            oid = rec["order_id"].strip()
            key = oid if oid else "__noorder__%d" % stats["rows"]
            if key in lib:
                for k, _ in FIELDS:
                    if not lib[key].get(k) and rec[k]:
                        lib[key][k] = rec[k]
                stats["merged"] += 1
            else:
                lib[key] = rec
                stats["added"] += 1
    return list(lib.values()), stats


def write_xlsx(records, out):
    records = sorted(records, key=sort_key)
    wb = Workbook()
    ws = wb.active
    ws.title = "开票汇总"
    ws.append([label for _, label in FIELDS])

    fill = PatternFill("solid", fgColor="2563EB")
    font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="D0D5DD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    for c in ws[1]:
        c.fill = fill
        c.font = font
        c.alignment = center
        c.border = border

    for rec in records:
        ws.append([rec.get(k, "") for k, _ in FIELDS])

    # 合计行
    total = sum(clean_amount(r.get("amount", "")) for r in records)
    last = [""] * len(FIELDS)
    last[0] = "合计 %d 单" % len(records)
    last[AMOUNT_IDX] = round(total, 2)
    ws.append(last)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.border = border
        if c.column == AMOUNT_IDX + 1:
            c.number_format = "0.00"

    widths = [20, 24, 22, 20, 14, 30, 18, 18, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(out)


def collect_inputs(inputs):
    files = []
    for inp in inputs:
        if os.path.isdir(inp):
            files += sorted(glob.glob(os.path.join(inp, "*.csv")))
            files += sorted(glob.glob(os.path.join(inp, "*.xlsx")))
            files += sorted(glob.glob(os.path.join(inp, "*.xlsm")))
        elif os.path.isfile(inp):
            files.append(inp)
        else:
            print(f"  [忽略] 不存在的路径: {inp}")
    return files


def main():
    ap = argparse.ArgumentParser(description="拼多多开票信息汇总器：合并多源 CSV/Excel 并按订单号去重，生成 xlsx 报表")
    ap.add_argument("-i", "--input", nargs="+", required=True, help="输入目录或文件（可多个）")
    ap.add_argument("-o", "--output", default="开票汇总.xlsx", help="输出 xlsx 路径（默认 ./开票汇总.xlsx）")
    args = ap.parse_args()

    files = collect_inputs(args.input)
    if not files:
        print("未找到任何可处理的 CSV/Excel 文件。")
        sys.exit(1)

    print(f"发现 {len(files)} 个输入文件：")
    for f in files:
        print("  - " + f)

    records, stats = aggregate(files)
    print(f"\n聚合结果：读取 {stats['rows']} 行 → 新增 {stats['added']} 条，合并补充 {stats['merged']} 条，去重后 {len(records)} 条。")

    out = args.output
    if not out.lower().endswith((".xlsx", ".xlsm")):
        out += ".xlsx"
    write_xlsx(records, out)
    total = sum(clean_amount(r.get("amount", "")) for r in records)
    print(f"已生成报表：{out}")
    print(f"金额合计：{total:,.2f} 元")


if __name__ == "__main__":
    main()
